"""Pruebas de api/chat_rutas.py (fase 2, 22/10) sin Supabase ni WordPress: db mockeado con monkeypatch."""
import io

import db
import api.chat_rutas as cr

CONV = {"id": "c1", "usuario": "franco", "titulo": "Gratitud y alegría", "mensajes": [
    {"role": "user", "content": "¿qué dice sobre la gratitud?"},
    {"role": "assistant", "content": "# Resumen\n\nLa gratitud es la **llave**.\n\n" + ("Texto largo. " * 60) +
                                     "\n\n| Tema | Fuente |\n|---|---|\n| Gratitud | Libro |\n",
     "mats": [{"tag": "CLIP", "titulo": "Clip", "texto": "x", "autor": "Br. David", "fuente": "youtube", "doc": "abc123def",
               "inicio_ms": 1000, "fin_ms": 5000}]},
    {"role": "user", "content": "gracias"},
    {"role": "assistant", "content": "De nada.", "mats": []},
]}


def _con_db(monkeypatch, conv=CONV):
    monkeypatch.setattr(db, "disponible", lambda: True)
    monkeypatch.setattr(db, "cargar_conversacion_de", lambda cid, u: dict(conv) if (cid == conv["id"] and u == conv["usuario"]) else None)


# ---------------- Renombrar, borrar, restaurar ----------------
def test_renombrar(logueado, monkeypatch):
    _con_db(monkeypatch)
    hecho = {}
    monkeypatch.setattr(db, "renombrar", lambda cid, t: hecho.update(cid=cid, titulo=t))
    r = logueado.patch("/conversaciones/c1", json={"titulo": "  Nuevo   título  "})
    assert r.status_code == 200 and r.json()["titulo"] == "Nuevo título" and hecho == {"cid": "c1", "titulo": "Nuevo título"}
    # vacío: queda el de antes
    assert logueado.patch("/conversaciones/c1", json={"titulo": "   "}).json()["titulo"] == "Gratitud y alegría"
    assert logueado.patch("/conversaciones/otra", json={"titulo": "x"}).status_code == 404


def test_borrar_y_restaurar(logueado, monkeypatch):
    _con_db(monkeypatch)
    borradas = []
    monkeypatch.setattr(db, "borrar_conversacion", lambda cid: borradas.append(cid))
    r = logueado.delete("/conversaciones/c1")
    assert r.status_code == 200 and borradas == ["c1"] and r.json()["dias"] == db.PAPELERA_DIAS
    assert logueado.delete("/conversaciones/otra").status_code == 404
    monkeypatch.setattr(db, "restaurar_conversacion_de", lambda cid, u: {"id": cid, "titulo": "T"} if cid == "c1" else None)
    assert logueado.post("/conversaciones/c1/restaurar").json() == {"ok": True, "id": "c1", "titulo": "T"}
    assert logueado.post("/conversaciones/zzz/restaurar").status_code == 404


def test_sin_supabase(logueado):
    assert logueado.patch("/conversaciones/c1", json={"titulo": "x"}).status_code == 503
    assert logueado.delete("/conversaciones/c1").status_code == 503
    assert logueado.post("/recortes", json={"video_id": "abc123def", "inicio_ms": 0, "fin_ms": 5000}).status_code == 503
    assert logueado.post("/publicaciones/borrador", json={"canal": "instagram", "texto": "x"}).status_code == 503


def test_sin_sesion(cliente):
    assert cliente.patch("/conversaciones/c1", json={"titulo": "x"}).status_code == 401
    assert cliente.get("/conversaciones/c1/mensajes/1/docx").status_code == 401
    assert cliente.post("/recortes", json={"video_id": "abc123def", "inicio_ms": 0, "fin_ms": 5000}).status_code == 401
    assert cliente.get("/banco").status_code == 401
    assert cliente.get("/publicaciones/modos").status_code == 401


# ---------------- Descargas ----------------
def test_docx_y_xlsx(logueado, monkeypatch):
    _con_db(monkeypatch)
    r = logueado.get("/conversaciones/c1/mensajes/1/docx")
    assert r.status_code == 200 and r.headers["content-type"].startswith(cr.TIPO_DOCX)
    cd = r.headers["content-disposition"]
    assert 'filename="Gratitud y alegria.docx"' in cd and "filename*=UTF-8''Gratitud%20y%20alegr%C3%ADa.docx" in cd
    from docx import Document
    doc = Document(io.BytesIO(r.content))
    assert doc.paragraphs[0].text == "Resumen" and len(doc.tables) == 1
    r = logueado.get("/conversaciones/c1/mensajes/1/xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith(cr.TIPO_XLSX)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Resumen"] and wb["Resumen"]["A2"].value == "Gratitud"
    # el mensaje 3 no tiene tablas; el 0 es de la persona; el 9 no existe
    assert logueado.get("/conversaciones/c1/mensajes/3/xlsx").status_code == 404
    assert logueado.get("/conversaciones/c1/mensajes/0/docx").status_code == 404
    assert logueado.get("/conversaciones/c1/mensajes/9/docx").status_code == 404


# ---------------- Recortes ----------------
def test_pedir_recorte(logueado, monkeypatch):
    monkeypatch.setattr(db, "disponible", lambda: True)
    pedidos = []

    def _pedir(vid, ini, fin, formato, titulo, usuario):
        pedidos.append((vid, ini, fin, formato, titulo, usuario))
        return "clip-1"
    monkeypatch.setattr(db, "pedir_clip", _pedir)
    r = logueado.post("/recortes", json={"video_id": "abc123def", "inicio_ms": 806000, "fin_ms": 850000, "formato": "vertical", "titulo": "La gratitud"})
    assert r.status_code == 200 and r.json()["id"] == "clip-1" and r.json()["estado"] == "pedido" and r.json()["ya"] is False
    assert pedidos == [("abc123def", 806000, 850000, "vertical", "La gratitud", "franco")]   # firmado por la persona de la sesión
    # validaciones del chat (nucleo): fin > inicio, menos de 15 minutos; id del video
    assert logueado.post("/recortes", json={"video_id": "abc123def", "inicio_ms": 5000, "fin_ms": 5000}).status_code == 400
    assert logueado.post("/recortes", json={"video_id": "abc123def", "inicio_ms": 0, "fin_ms": 16 * 60 * 1000}).status_code == 400
    assert logueado.post("/recortes", json={"video_id": "a b", "inicio_ms": 0, "fin_ms": 5000}).status_code == 400
    # ya pedido: no es error, cuenta el estado que tiene
    monkeypatch.setattr(db, "pedir_clip", lambda *a: "ya")
    monkeypatch.setattr(db, "clips_de", lambda vids: [{"id": "c9", "video_id": "abc123def", "inicio_ms": 806000, "fin_ms": 850000,
                                                       "formato": "horizontal", "estado": "listo", "url": "https://x/clip.mp4", "error": None}])
    r = logueado.post("/recortes", json={"video_id": "abc123def", "inicio_ms": 806000, "fin_ms": 850000, "formato": "horizontal"})
    assert r.status_code == 200 and r.json()["ya"] is True and r.json()["estado"] == "listo" and r.json()["url"] == "https://x/clip.mp4"


def test_recortes_de_la_conversacion(logueado, monkeypatch):
    _con_db(monkeypatch)
    visto = []
    monkeypatch.setattr(db, "clips_de", lambda vids: visto.append(vids) or [{"id": "c1", "video_id": "abc123def", "inicio_ms": 1000,
                                                                              "fin_ms": 5000, "formato": "vertical", "estado": "pedido", "url": None, "error": None}])
    r = logueado.get("/conversaciones/c1/recortes")
    assert r.status_code == 200 and len(r.json()) == 1 and visto == [["abc123def"]]


# ---------------- Pieza propuesta ----------------
def test_modos_y_banco(logueado, monkeypatch):
    monkeypatch.setattr(cr, "modos", lambda: {"wordpress": "borrador", "instagram": "real", "facebook": "apagado", "youtube": "privado"})
    d = logueado.get("/publicaciones/modos").json()
    assert d["canales"]["instagram"]["a_aprobar"] is True and d["canales"]["facebook"]["a_aprobar"] is False
    assert d["canales"]["youtube"]["a_aprobar"] is False and d["canales"]["wordpress"]["modo"] == "borrador"
    assert d["publica"] is True and "franco" in d["aprobadores"]
    # sin WordPress configurado, el banco devuelve lista vacía (no error)
    assert logueado.get("/banco?consulta=amanecer").json() == []


def test_guardar_borrador(logueado, monkeypatch):
    monkeypatch.setattr(db, "disponible", lambda: True)
    guardado = {}

    def _guardar(usuario, canal, titulo, texto, imagen_url, fecha, pieza, pid=None):
        guardado.update(usuario=usuario, canal=canal, titulo=titulo, texto=texto, imagen_url=imagen_url, fecha=fecha, pieza=pieza, pid=pid)
        return 77
    monkeypatch.setattr(db, "guardar_borrador", _guardar)
    r = logueado.post("/publicaciones/borrador", json={"canal": "instagram", "titulo": "Gratitud", "texto": "La gratitud es la llave.",
                                                       "imagen_url": "https://wp/foto.jpg", "fecha": "2026-10-30", "hora": "9:30", "imagen_busqueda": "amanecer"})
    assert r.status_code == 200 and r.json()["id"] == 77 and r.json()["estado"] == "borrador"
    assert guardado["canal"] == "instagram" and guardado["fecha"] == "2026-10-30T09:30:00-03:00" and guardado["pid"] is None
    assert guardado["pieza"]["busq"] == "amanecer" and guardado["pieza"]["fecha"] == "2026-10-30" and guardado["pieza"]["hora"] == "09:30"
    assert guardado["pieza"]["url"] == "https://wp/foto.jpg" and guardado["pieza"]["origen"] == "web"
    assert logueado.post("/publicaciones/borrador", json={"canal": "", "texto": "x"}).status_code == 400
    assert logueado.post("/publicaciones/borrador", json={"canal": "instagram", "texto": "x", "fecha": "30/10/2026"}).status_code == 400


def test_enviar_a_aprobacion(logueado, monkeypatch):
    monkeypatch.setattr(db, "disponible", lambda: True)
    monkeypatch.setattr(cr, "modos", lambda: {"wordpress": "borrador", "instagram": "real", "facebook": "apagado", "youtube": "privado"})
    propuesto = {}

    def _proponer(usuario, canal, modo, titulo, texto, imagen_url, fecha, extra=None):
        propuesto.update(usuario=usuario, canal=canal, modo=modo, titulo=titulo, texto=texto, imagen_url=imagen_url, fecha=fecha, extra=extra)
        return 88
    monkeypatch.setattr(db, "proponer_publicacion", _proponer)
    borrados = []
    monkeypatch.setattr(db, "_req", lambda *a, **k: borrados.append((a, k)))
    r = logueado.post("/publicaciones/a-aprobar", json={"canal": "instagram", "texto": "**Gratitud** es la [llave](https://x.y).", "borrador_id": "5a1e"})
    assert r.status_code == 200, r.text
    assert r.json()["id"] == 88 and r.json()["estado"] == "a_aprobar" and "franco" in r.json()["mensaje"]
    assert propuesto["modo"] == "real" and propuesto["texto"] == "Gratitud es la llave." and propuesto["fecha"] is None
    assert propuesto["extra"]["origen"] == "web" and propuesto["imagen_url"] is None
    assert any(b[0][0] == "DELETE" and "estado=eq.borrador" in b[0][1] for b in borrados)   # el borrador retomado se elimina (roles() también consulta la tabla usuarios)
    # WordPress conserva el markdown/HTML y exige título
    assert logueado.post("/publicaciones/a-aprobar", json={"canal": "wordpress", "texto": "x"}).status_code == 400
    r = logueado.post("/publicaciones/a-aprobar", json={"canal": "wordpress", "titulo": "T", "texto": "**x**"})
    assert r.status_code == 200 and propuesto["texto"] == "**x**" and propuesto["modo"] == "borrador"
    # canal apagado, YouTube, sin texto, sin canal
    assert logueado.post("/publicaciones/a-aprobar", json={"canal": "facebook", "texto": "x"}).status_code == 400
    assert logueado.post("/publicaciones/a-aprobar", json={"canal": "youtube", "titulo": "T", "texto": "x"}).status_code == 400
    assert logueado.post("/publicaciones/a-aprobar", json={"canal": "instagram", "texto": "  "}).status_code == 400
    assert logueado.post("/publicaciones/a-aprobar", json={"canal": "newsletter", "texto": "x"}).status_code == 400

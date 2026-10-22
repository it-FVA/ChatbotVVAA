"""documentos.py — una respuesta del asistente (markdown simple) como Word o Excel. Compartido por Streamlit (app.py) y la API
(api/chat_rutas.py: GET /conversaciones/{id}/mensajes/{n}/docx|xlsx). (22/10) Movido desde app.py sin cambiar la lógica:
md_a_docx (20/10, Pato pidió el concepto de campaña "en un Word"), tablas_md y tablas_a_xlsx (21/10, Julián: "que también se
pueda en Excel").

  md_a_docx(texto, titulo) -> bytes .docx      títulos #, listas -, * y 1., citas >, negritas **, tablas | … |
  tablas_md(texto) -> [(titulo_cercano, filas)]  las tablas markdown de la respuesta, celdas limpias
  tablas_a_xlsx(tablas, titulo) -> bytes .xlsx  una hoja por tabla
  nombre_archivo(titulo) -> str                 nombre seguro para el archivo (sin extensión)

Umbral de Streamlit para ofrecer «⬇️ Word»: respuestas de 600 caracteres o más (LARGO_MINIMO_WORD).
"""
LARGO_MINIMO_WORD = 600


def nombre_archivo(titulo):
    """Como app.descargar_en_chat: solo letras, números, espacio, guion y guion bajo; 60 caracteres."""
    t = (titulo or "Documento del asistente").strip()
    return "".join(c if c.isalnum() or c in " -_" else "" for c in t).strip()[:60] or "documento"


def md_a_docx(texto, titulo="Documento"):
    """(20/10, Pato pidió el concepto de campaña "en un Word") Convierte la respuesta (markdown simple) a .docx:
    títulos #, listas -, * y 1., citas >, negritas **. Lo demás va como párrafo."""
    import io, re
    from docx import Document
    doc = Document()
    doc.core_properties.title = titulo[:200]

    def _runs(par, linea):
        for i, trozo in enumerate(re.split(r"(\*\*[^*]+\*\*)", linea)):
            if not trozo:
                continue
            if trozo.startswith("**") and trozo.endswith("**"):
                par.add_run(trozo[2:-2]).bold = True
            else:
                par.add_run(re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", trozo).replace("*", ""))

    lineas = (texto or "").splitlines()
    k = 0
    while k < len(lineas):
        s = lineas[k].rstrip()
        k += 1
        if not s.strip():
            continue
        if s.strip() in ("---", "***"):
            continue
        if s.strip().startswith("|") and s.strip().endswith("|"):     # (21/10) tabla markdown → tabla de Word
            bloque = [s]
            while k < len(lineas) and lineas[k].strip().startswith("|"):
                bloque.append(lineas[k]); k += 1
            tablas = tablas_md("\n".join(bloque))
            for _t, filas in tablas:
                ancho = max(len(f) for f in filas)
                tabla = doc.add_table(rows=len(filas), cols=ancho)
                tabla.style = "Table Grid"
                for r, fila in enumerate(filas):
                    for c in range(ancho):
                        celda = tabla.cell(r, c); celda.text = ""
                        run = celda.paragraphs[0].add_run(fila[c] if c < len(fila) else "")
                        if r == 0:
                            run.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)), 4))
            continue
        if re.match(r"^\s*[-*•]\s+", s):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", s))
            continue
        if re.match(r"^\s*\d+[.)]\s+", s):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", s))
            continue
        if s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            _runs(p, s.lstrip("> ").strip())
            continue
        _runs(doc.add_paragraph(), s)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def tablas_md(texto):
    """(21/10, Julián: "me lo deja descargar solo como Word; que también se pueda en Excel") Saca las tablas
    markdown de una respuesta: lista de (titulo_cercano, filas) con filas = listas de celdas ya limpias
    (<br> → salto de línea, sin ** ni *)."""
    import re
    tablas, filas, titulo, ultimo_titulo = [], [], "", ""
    for linea in (texto or "").splitlines() + [""]:
        s = linea.strip()
        if s.startswith("|") and s.endswith("|"):
            celdas = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in celdas if c):   # la línea |---|---| separadora
                continue
            celdas = [re.sub(r"<br\s*/?>", "\n", c, flags=re.I) for c in celdas]
            celdas = [re.sub(r"\*\*([^*]+)\*\*", r"\1", c).replace("*", "").strip() for c in celdas]
            if not filas:
                titulo = ultimo_titulo
            filas.append(celdas)
        else:
            if filas:
                tablas.append((titulo, filas)); filas = []
            mt = re.match(r"^#{1,4}\s+(.*)", s)
            if mt:
                ultimo_titulo = mt.group(1).replace("**", "").strip()
    return tablas


def tablas_a_xlsx(tablas, titulo_doc="Planilla"):
    """Una hoja por tabla: encabezado en negrita, texto con ajuste de línea, anchos razonables."""
    import io, re
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = Workbook(); wb.remove(wb.active)
    usados = set()
    for n, (titulo, filas) in enumerate(tablas, 1):
        nombre = re.sub(r"[\[\]:*?/\\]", "", titulo or f"Tabla {n}")[:28] or f"Tabla {n}"
        if nombre in usados:
            nombre = f"{nombre[:25]} {n}"
        usados.add(nombre)
        ws = wb.create_sheet(nombre)
        ancho = max(len(f) for f in filas)
        for r, fila in enumerate(filas, 1):
            for c in range(ancho):
                celda = ws.cell(row=r, column=c + 1, value=(fila[c] if c < len(fila) else ""))
                celda.alignment = Alignment(wrap_text=True, vertical="top")
                if r == 1:
                    celda.font = Font(bold=True); celda.fill = PatternFill("solid", fgColor="FDE9C9")
        for c in range(ancho):
            largo = max((len(x) for f in filas for x in [f[c] if c < len(f) else ""]), default=10)
            ws.column_dimensions[ws.cell(row=1, column=c + 1).column_letter].width = min(60, max(14, largo // 3))
        ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
